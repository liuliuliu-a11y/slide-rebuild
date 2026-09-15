"""Rebuild this fixed demonstration layout, not an automatic image converter.

Requires python-pptx and Pillow. Run with --out to choose the PPTX location.
"""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
NAVY, TEAL, AMBER, WHITE, MUTED = '16324F', '078C95', 'EF931D', 'FFFFFF', '506880'
SCALE = 1672 / (40 / 3)


def unit(px):
    return Inches(px / SCALE)


def color(value):
    return RGBColor.from_string(value)


def font(run, size, rgb=NAVY, bold=False):
    run.font.name = 'Microsoft YaHei'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(rgb)
    east = OxmlElement('a:ea')
    east.set('typeface', 'Microsoft YaHei')
    run._r.get_or_add_rPr().append(east)


def write(shape, text, size=16, rgb=NAVY, bold=False, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = unit(3)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    font(run, size, rgb, bold)


def text(slide, name, value, box, size=16, rgb=NAVY, bold=False, center=False):
    sh = slide.shapes.add_textbox(*(unit(v) for v in box))
    sh.name = name
    write(sh, value, size, rgb, bold, PP_ALIGN.CENTER if center else PP_ALIGN.LEFT)
    return sh


def rect(slide, name, box, fill=WHITE, stroke=None, rounded=False):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE, *(unit(v) for v in box))
    sh.name = name
    style = sh._element.find(qn('p:style'))
    if style is not None: sh._element.remove(style)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color(fill)
    if stroke:
        sh.line.color.rgb = color(stroke)
        sh.line.width = Pt(1.4)
    else:
        sh.line.fill.background()
    if rounded:
        sh.adjustments[0] = 0.12
    return sh


def connector(slide, name, start, end, rgb=TEAL, width=2, dashed=False, begin=None, finish=None, arrowhead=True):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, unit(start[0]), unit(start[1]), unit(end[0]), unit(end[1]))
    sh.name = name
    style = sh._element.find(qn('p:style'))
    if style is not None: sh._element.remove(style)
    sh.line.color.rgb = color(rgb)
    sh.line.width = Pt(width)
    ln = sh.line._get_or_add_ln()
    if dashed:
        dash = OxmlElement('a:prstDash'); dash.set('val', 'dash'); ln.append(dash)
    if arrowhead:
        arrow = OxmlElement('a:tailEnd'); arrow.set('type', 'triangle'); ln.append(arrow)
    if begin:
        sh.begin_connect(begin[0], begin[1])
    if finish:
        sh.end_connect(finish[0], finish[1])
    return sh


def build(out):
    d = json.loads((HERE / 'data.json').read_text(encoding='utf-8'))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(40/3), Inches(7.5)
    prs.core_properties.author = 'Slide Rebuild contributors'
    prs.core_properties.last_modified_by = 'Slide Rebuild contributors'
    prs.core_properties.title = '工业储热系统：公开复刻示例'
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = color(WHITE)
    text(slide, 'title', '工业储热系统', (28, 7, 480, 91), 43, bold=True)
    text(slide, 'subtitle', '能量流与运行状态', (540, 20, 650, 64), 25, MUTED, True)
    badge = rect(slide, 'snapshot', (1340, 24, 304, 58), NAVY, rounded=True)
    write(badge, '12:00  |  虚构运行快照', 14, WHITE, align=PP_ALIGN.CENTER)
    for i, (label, value) in enumerate([
        ('额定储能量', f"{d['capacity_kwh']} kWh"), ('充电功率上限', f"{d['charge_limit_kw']} kW"),
        ('示例放热功率', f"{d['heat_output_kw']} kW"), ('当前 SOC', f"{d['soc']:.0%}")]):
        x = 22 + i * 418
        a = text(slide, f'metric-{i}-label', label, (x, 113, 399, 37), 15, MUTED, center=True)
        b = text(slide, f'metric-{i}-value', value, (x, 151, 399, 65), 34, TEAL, True, True)
        group = slide.shapes.add_group_shape([a, b]); group.name = f'metric-{i}'
    rect(slide, 'left-panel', (22, 231, 956, 598), WHITE, 'DCE5EB')
    rect(slide, 'chart-panel', (988, 231, 663, 343), WHITE, 'DCE5EB')
    for name, label, box in [('diagram-header','能量路径与调度',(22,231,956,54)), ('chart-header','SOC 示例曲线',(988,231,663,54)), ('table-header','计算口径',(988,584,663,54))]:
        sh=rect(slide,name,box,NAVY); write(sh,label,21,WHITE,True)
        sh.text_frame.margin_left=unit(29)
        rect(slide,name+'-accent',(box[0],box[1],13,box[3]),TEAL)
    equipment=slide.shapes.add_picture(str(HERE/'equipment.png'),unit(388),unit(353),width=unit(392),height=unit(392))
    equipment.name='equipment-image'
    text(slide,'equipment-label','储热单元',(457,730,255,38),15,bold=True,center=True)
    nodes={}
    for name,label,box in [('grid-node','电网',(44,494,123,64)),('heater-node','电加热',(243,496,113,61)),('user-node','热用户',(838,495,117,65)),('control-node','调度控制',(370,302,160,47))]:
        sh=rect(slide,name,box,NAVY if name=='control-node' else WHITE,NAVY,True)
        write(sh,label,15,WHITE if name=='control-node' else NAVY,True,PP_ALIGN.CENTER);nodes[name]=sh
    connector(slide,'grid-to-heater',(170,526),(241,526),AMBER,3,begin=(nodes['grid-node'],3),finish=(nodes['heater-node'],1))
    ports={}
    for name,box,rgb in [('storage-input-port',(422,522,7,7),AMBER),('storage-output-port',(754,522,7,7),TEAL),('storage-control-port',(446,410,8,8),NAVY)]:
        ports[name]=rect(slide,name,box,rgb,rounded=True)
    connector(slide,'heater-to-storage',(358,526),(422,526),AMBER,3,begin=(nodes['heater-node'],3),finish=(ports['storage-input-port'],1))
    connector(slide,'storage-to-user',(761,526),(835,526),TEAL,3,begin=(ports['storage-output-port'],3),finish=(nodes['user-node'],1))
    connector(slide,'control-to-heater-route',(370,325),(300,325),NAVY,1.5,True,begin=(nodes['control-node'],1),arrowhead=False)
    connector(slide,'control-to-heater',(300,325),(300,496),NAVY,1.5,True,finish=(nodes['heater-node'],0))
    connector(slide,'control-to-storage',(450,349),(450,410),NAVY,1.5,True,begin=(nodes['control-node'],2),finish=(ports['storage-control-port'],0))
    connector(slide,'legend-energy',(48,799),(120,799),AMBER,2.5)
    text(slide,'legend-energy-text','实线：能量流',(139,779,215,36),13)
    connector(slide,'legend-control',(347,799),(413,799),NAVY,1.4,True)
    text(slide,'legend-control-text','虚线：控制指令',(431,779,300,36),13)
    graphic=slide.shapes.add_picture(str(HERE/'soc-curve.png'),unit(995),unit(296),width=unit(650),height=unit(273))
    graphic.name='soc-matlab-plot'
    table=slide.shapes.add_table(4,2,unit(988),unit(638),unit(663),unit(191));table.name='assumptions-table'
    table.table.columns[0].width=unit(315);table.table.columns[1].width=unit(348)
    rows=[['项目','约定'],['功率与 SOC','仅为演示数据'],['储能量计算','按额定值与 SOC'],['净供热量','本页不计算']]
    for ri,row in enumerate(rows):
        for ci,value in enumerate(row):
            cell=table.table.cell(ri,ci);cell.text=value;cell.margin_left=unit(29);cell.margin_top=0;cell.margin_bottom=0;cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=color('EAF0F4' if ri==0 else WHITE)
            tcpr=cell._tc.get_or_add_tcPr()
            for edge in ['lnL','lnR','lnT','lnB']:
                border=OxmlElement('a:'+edge);border.set('w','6350')
                fill=OxmlElement('a:solidFill');rgb=OxmlElement('a:srgbClr');rgb.set('val','DCE5EB');fill.append(rgb);border.append(fill);tcpr.append(border)
            for para in cell.text_frame.paragraphs:
                for run in para.runs:font(run,13,NAVY,ri==0)
    label=rect(slide,'energy-label',(22,850,253,69),NAVY);write(label,'账面储能量',23,WHITE,True,PP_ALIGN.CENTER)
    e=d['capacity_kwh']*d['soc']
    equation=text(slide,'formula',f"E = {d['capacity_kwh']} × {d['soc']:.0%} = {e:g} kWh",(312,855,680,57),24,bold=True)
    para=equation.text_frame.paragraphs[0];para.clear()
    for value,rgb in [(f"E = {d['capacity_kwh']} × {d['soc']:.0%} = ",NAVY),(f'{e:g}',AMBER),(' kWh',NAVY)]:
        run=para.add_run();run.text=value;font(run,24,rgb,True)
    text(slide,'disclaimer','公开复刻示例｜参数与曲线均为虚构；设备图为示意，不代表实际产品或可交付净热量。',(1004,891,640,35),9,MUTED)
    slide.notes_slide.notes_text_frame.text='公开演示。所有参数和曲线均为虚构；设备图为示意。SOC 曲线由 MATLAB 数据与脚本再生成，在 PPT 中为独立图片。公式是可编辑文本，不是 Office 公式对象。'
    out.parent.mkdir(parents=True,exist_ok=True);prs.save(out)
    print(out)


if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--out',type=Path,default=HERE/'rebuilt.pptx')
    build(ap.parse_args().out)
